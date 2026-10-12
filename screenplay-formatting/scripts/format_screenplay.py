#!/usr/bin/env python3
"""Reference Arabic screenplay formatter. No API keys or remote uploads."""
import argparse
import hashlib
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import uuid
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as ET

from PIL import ImageFont, features
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fontTools.ttLib import TTFont
from screenplay_io import FIELDS, load_source, validate_source

ROOT = Path(__file__).resolve().parents[1]
MM_PT = 72 / 25.4
DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def arabic_number(value):
    return str(value).translate(DIGITS)


def words(value):
    return value.split()


@dataclass
class Style:
    profile: str = "reference"
    family: str = "Amiri"
    regular: str = str(ROOT / "assets/fonts/Amiri-Regular.ttf")
    bold: str = str(ROOT / "assets/fonts/Amiri-Bold.ttf")
    body: float = 16
    line: float = 26
    gap: float = 8
    location: float = 20
    meta: float = 14
    margin: float = 20
    top: float = 25
    bottom: float = 20
    columns: tuple = (66, 22, 8, 74)
    safety: float = 18

    def validate(self):
        if not 12 <= self.body <= 24 or self.line < self.body * 1.45:
            raise ValueError("استخدم متن ١٢–٢٤ نقطة، وارتفاع سطر لا يقل عن ١٫٤٥ من حجم الخط.")
        for path in (self.regular, self.bold):
            if not Path(path).is_file():
                raise ValueError(f"الخط غير موجود: {path}")
            font = TTFont(path)
            families = {n.toUnicode().lower() for n in font["name"].names if n.nameID in (1, 16)}
            if self.family.lower() not in families:
                raise ValueError("اسم عائلة الخط لا يطابق ملف الخط.")
            if "OS/2" in font and font["OS/2"].fsType & 0x0002:
                raise ValueError("الخط يمنع التضمين؛ استخدم خطاً يسمح بتضمينه.")
            font.close()
        if not features.check_feature("raqm"):
            raise RuntimeError("Pillow يحتاج دعم RAQM لتوزيع أسطر عربية موثوق.")
        return self


class Measure:
    def __init__(self, style):
        self.style = style
        self.cache = {}

    def font(self, size, bold=False):
        key = (size, bold)
        if key not in self.cache:
            self.cache[key] = ImageFont.truetype(self.style.bold if bold else self.style.regular,
                                               round(size * 4), layout_engine=ImageFont.Layout.RAQM)
        return self.cache[key]

    def wrap(self, text, width_mm, size=None, bold=False):
        if not text.strip():
            return []
        font = self.font(size or self.style.body, bold)
        width = width_mm * MM_PT * 4 - 4
        output = []
        for paragraph in text.splitlines() or [text]:
            if not paragraph.strip():
                output.append("")
                continue
            current = ""
            for word in paragraph.split():
                candidate = f"{current} {word}" if current else word
                if font.getlength(candidate, direction="rtl") <= width:
                    current = candidate
                else:
                    if current:
                        output.append(current)
                    if font.getlength(word, direction="rtl") > width:
                        raise ValueError(f"كلمة لا تسع عمودها بالحجم المحدد: {word[:70]}")
                    current = word
            if current:
                output.append(current)
        return output


def paginate(source, style):
    source = validate_source(source)
    m = Measure(style)
    pages = []
    number = source["first_page_number"]
    capacity = (297 - style.top - style.bottom) * MM_PT - style.safety

    def add_page(episode, scene=None, first=False, cover=False):
        nonlocal number
        page = {"number": number, "episode": episode["number"], "episode_label": episode["label"],
                "cover": cover, "scene": scene["number"] if scene else "", "first_scene_page": first,
                "heading": (scene.get("location") or scene.get("heading") or "") if scene else "",
                "interior": scene.get("interior", "") if scene else "",
                "time": scene.get("time", "") if scene else "", "items": [], "used": 0.0}
        if first:
            meta_lines = max(1, len(m.wrap(f"المشهد ({arabic_number(scene['number'])})", 85, style.meta)))
            location_lines = len(m.wrap(page["heading"], 170, style.location, True))
            page["used"] = meta_lines * 23 + location_lines * (style.location * 1.55) + 18
        pages.append(page)
        number += 1
        return page

    for ei, episode in enumerate(source["episodes"]):
        add_page(episode, cover=True)
        for si, scene in enumerate(episode["scenes"]):
            page = add_page(episode, scene, first=True)
            for bi, beat in enumerate(scene["beats"]):
                uid = f"{ei}.{si}.{bi}"
                if beat.get("transition", "").strip():
                    lines = m.wrap(beat["transition"], 90, 18, True)
                    height = len(lines) * 28 + 10
                    if height > capacity:
                        raise ValueError("نص الانتقال يتجاوز صفحة كاملة.")
                    if page["used"] + height > capacity:
                        # Keep the end of the preceding beat beside its transition.
                        if page["items"] and page["items"][-1]["height"] + height <= capacity:
                            tail = page["items"].pop()
                            page["used"] -= tail["height"]
                            page = add_page(episode, scene)
                            page["items"].append(tail)
                            page["used"] += tail["height"]
                        else:
                            page = add_page(episode, scene)
                    page["items"].append({"kind": "transition", "uid": uid, "fields": {"transition": lines}, "height": height})
                    page["used"] += height
                    continue
                field_lines = {}
                for key, width in (("dialogue", style.columns[0]), ("character", style.columns[1]),
                                   ("action", style.columns[3] - 4), ("parenthetical", style.columns[3] - 4)):
                    field_lines[key] = m.wrap(beat.get(key, ""), width)
                action = [(key, line) for key in ("action", "parenthetical") for line in field_lines[key]]
                spoken = field_lines["dialogue"]
                names = field_lines["character"]
                count = max(len(action), len(spoken), len(names), 1)
                full_height = count * style.line + style.gap
                # Keep short replies intact; split longer speeches into the remaining space.
                if count <= 4 and full_height <= capacity and page["used"] + full_height > capacity and page["items"]:
                    page = add_page(episode, scene)
                offset = 0
                while offset < count:
                    available = math.floor((capacity - page["used"] - style.gap) / style.line)
                    minimum = max(len(names), min(2, len(spoken))) if offset == 0 else 1
                    if available < max(1, minimum):
                        page = add_page(episode, scene)
                        available = math.floor((capacity - style.gap) / style.line)
                    take = min(available, count - offset)
                    if count - offset - take == 1 and take > minimum:
                        take -= 1
                    selected = {key: [] for key in ("action", "parenthetical", "dialogue", "character")}
                    for key, line in action[offset:offset + take]:
                        selected[key].append(line)
                    selected["dialogue"] = spoken[offset:offset + take]
                    selected["character"] = names if offset == 0 else []
                    height = take * style.line + style.gap
                    page["items"].append({"kind": "beat", "uid": uid, "continuation": offset > 0,
                                          "fields": selected, "height": height})
                    page["used"] += height
                    offset += take
                    if offset < count:
                        page = add_page(episode, scene)
    reconstructed = defaultdict(lambda: defaultdict(list))
    for page in pages:
        for item in page["items"]:
            for key, lines in item["fields"].items():
                reconstructed[item["uid"]][key].extend(lines)
    for ei, episode in enumerate(source["episodes"]):
        for si, scene in enumerate(episode["scenes"]):
            for bi, beat in enumerate(scene["beats"]):
                uid = f"{ei}.{si}.{bi}"
                for key in FIELDS:
                    if words(beat.get(key, "")) != words(" ".join(reconstructed[uid][key])):
                        raise RuntimeError(f"فشل حفظ نص {uid}.{key}")
    return pages


def xml_child(parent, name, **attrs):
    element = OxmlElement(f"w:{name}")
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), str(value))
    parent.append(element)
    return element


def font_run(run, style, size, bold=False, rtl=True):
    run.font.name = style.family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    prop = run._element.get_or_add_rPr()
    fonts = prop.find(qn("w:rFonts"))
    if fonts is None:
        fonts = xml_child(prop, "rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), style.family)
    xml_child(prop, "szCs", val=round(size * 2))
    xml_child(prop, "bCs", val="1" if bold else "0")
    xml_child(prop, "rtl", val="1" if rtl else "0")
    xml_child(prop, "lang", val="ar-KW", bidi="ar-KW")


def paragraph(container, content, style, size=None, bold=False, align="right", line=None, after=0, before=0):
    p = container.add_paragraph()
    set_paragraph(p, content, style, size, bold, align, line, after, before)
    return p


def set_paragraph(p, content, style, size=None, bold=False, align="right", line=None, after=0, before=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(line or style.line)
    p.paragraph_format.widow_control = True
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    xml_child(p._p.get_or_add_pPr(), "bidi", val="1")
    # start/end are logical alignment values: start is the right edge in RTL.
    if align in ("right", "left"):
        p._p.get_or_add_pPr().find(qn("w:jc")).set(qn("w:val"), "start" if align == "right" else "end")
    run = p.add_run(content)
    font_run(run, style, size or style.body, bold)
    return p


def remove_table_borders(table):
    props = table._tbl.tblPr
    xml_child(props, "bidiVisual", val="0")
    borders = xml_child(props, "tblBorders")
    for key in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml_child(borders, key, val="nil")


def table(parent, widths, style):
    if not isinstance(parent, DocxDocument):
        t = parent.add_table(rows=1, cols=len(widths), width=Mm(sum(widths)))
    else:
        t = parent.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(t)
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    for cell, width in zip(t.rows[0].cells, widths):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        margin = xml_child(cell._tc.get_or_add_tcPr(), "tcMar")
        for key in ("top", "left", "bottom", "right"):
            xml_child(margin, key, w="0", type="dxa")
        p = cell.paragraphs[0]
        set_paragraph(p, "", style, size=1, line=1)
    return t


def cell_text(cell, content, style, size=None, bold=False, align="right", line=None, after=0):
    p = cell.paragraphs[0]
    p.clear()
    set_paragraph(p, content, style, size, bold, align, line, after)


def blank(doc, height):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(max(1, height))
    p.add_run().font.size = Pt(1)
    return p


def rule_cell(cell, double=False):
    borders = xml_child(cell._tc.get_or_add_tcPr(), "tcBorders")
    xml_child(borders, "bottom", val="double" if double else "single", sz="6", color="000000")


def rules(section, page, source, style):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    footer = section.footer
    set_paragraph(header.paragraphs[0], "", style, size=1, line=1)
    set_paragraph(footer.paragraphs[0], "", style, size=1, line=1)
    odd = page["number"] % 2 == 1
    if not page["cover"]:
        widths = (14, 30, 126) if odd else (126, 30, 14)
        t = table(header, widths, style)
        label = source["part_label"] if odd else source["title"]
        # Do not invent a part when metadata was not provided.
        if not label:
            label = source["title"]
        cell_text(t.cell(0, 1), label, style, size=12, align="center", line=19)
        rule_cell(t.cell(0, 0), True)
        rule_cell(t.cell(0, 2), True)
    t = table(footer, (14, 16, 140) if odd else (140, 16, 14), style)
    cell_text(t.cell(0, 1), arabic_number(page["number"]), style, size=15, bold=True, align="center", line=22)
    rule_cell(t.cell(0, 0))
    rule_cell(t.cell(0, 2))


def configure_section(section, cover=False):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = Mm(15 if cover else 25)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def cover(doc, page, source, style, m):
    p = source["production"]
    t = table(doc, (85, 85), style)
    cell_text(t.cell(0, 0), p["sector"], style, size=17, bold=True, align="left", line=28)
    cell_text(t.cell(0, 1), p["organization"], style, size=17, bold=True, line=28)
    paragraph(doc, p["unit"], style, size=18, bold=True, align="center", line=30, before=8)
    blank(doc, 108)
    title = "\n".join(m.wrap(source["title"], 170, 44, True))
    paragraph(doc, title, style, size=44, bold=True, align="center", line=62)
    paragraph(doc, source["part_label"], style, size=24, align="center", line=40, after=74)
    t = table(doc, (85, 85), style)
    cell_text(t.cell(0, 0), source["director"], style, size=20, bold=True, align="center", line=34)
    cell_text(t.cell(0, 1), source["author"], style, size=20, bold=True, align="center", line=34)
    blank(doc, 8)
    t = table(doc, (85, 85), style)
    cell_text(t.cell(0, 0), f"E : {page['episode']}", style, size=34, align="center", line=52)
    xml_child(t.cell(0, 0).paragraphs[0]._p.get_or_add_pPr(), "bidi", val="0")
    for run in t.cell(0, 0).paragraphs[0].runs:
        xml_child(run._element.get_or_add_rPr(), "rtl", val="0")
    cell_text(t.cell(0, 1), f"ح : {arabic_number(page['episode'])}", style, size=30, align="center", line=52)
    blank(doc, 22)
    paragraph(doc, page["episode_label"] or f"الحلقة {arabic_number(page['episode'])}",
              style, size=23, bold=True, align="center", line=40)


def render_docx(source, pages, style, output):
    doc = Document()
    doc.core_properties.title = source["title"]
    doc.core_properties.author = source["author"]
    normal = doc.styles["Normal"]
    normal.font.name = style.family
    normal.font.size = Pt(style.body)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    settings = doc.settings.element
    # Keep columns physically LTR and shape each Arabic paragraph individually.
    xml_child(settings, "themeFontLang", val="ar-KW", bidi="ar-KW")
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = xml_child(settings, "compat")
    xml_child(compat, "doNotExpandShiftReturn", val="1")
    m = Measure(style)
    text_checks = []
    for index, page in enumerate(pages):
        if index:
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            end = doc.paragraphs[-1]
            end.paragraph_format.line_spacing = Pt(1)
            end.paragraph_format.space_before = end.paragraph_format.space_after = Pt(0)
            end.add_run().font.size = Pt(1)
        else:
            section = doc.sections[0]
        configure_section(section, page["cover"])
        rules(section, page, source, style)
        if page["cover"]:
            cover(doc, page, source, style, m)
            continue
        if page["first_scene_page"]:
            t = table(doc, (85, 85), style)
            setting = " / ".join(x for x in (page["interior"], page["time"]) if x)
            cell_text(t.cell(0, 0), setting, style, size=style.meta, align="left", line=23)
            cell_text(t.cell(0, 1), f"المشهد ({arabic_number(page['scene'])})" if page["scene"] else "",
                      style, size=style.meta, line=23)
            if page["heading"]:
                paragraph(doc, "\n".join(m.wrap(page["heading"], 170, style.location, True)),
                          style, size=style.location, bold=True, align="center", line=style.location * 1.55)
            blank(doc, 18)
        for item in page["items"]:
            fields = item["fields"]
            if item["kind"] == "transition":
                label = "\n".join(fields["transition"])
                if label.strip() == "قطع":
                    label = "ق" + "\u0640" * 18 + "طع"
                p = paragraph(doc, "- " + label + " -",
                              style, size=18, bold=True, align="center", line=28, after=10)
                continue
            t = table(doc, style.columns, style)
            row = t.rows[0]
            row.height = Pt(item["height"])
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            xml_child(row._tr.get_or_add_trPr(), "cantSplit")
            dialogue = "\n".join(fields.get("dialogue", []))
            name = "\n".join(fields.get("character", []))
            action = "\n".join(fields.get("action", []) + fields.get("parenthetical", []))
            if action and not item["continuation"]:
                action = "- " + action
            cell_text(t.cell(0, 0), dialogue, style, align="justify")
            cell_text(t.cell(0, 1), name, style)
            cell_text(t.cell(0, 2), "", style, size=1, line=1)
            cell_text(t.cell(0, 3), action, style, align="right")
            for cell, expected in ((t.cell(0, 0), dialogue), (t.cell(0, 1), name), (t.cell(0, 3), action)):
                text_checks.append((cell, expected))
    for cell, expected in text_checks:
        if words(cell.text) != words(expected):
            raise RuntimeError("فشل حفظ نص الفقرة في Word.")
    doc.save(output)
    embed_fonts(output, style)
    saved = Document(output)
    actual_cells = [cell.text for t in saved.tables for row in t.rows for cell in row.cells]
    expected_cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    if actual_cells != expected_cells or [p.text for p in saved.paragraphs] != [p.text for p in doc.paragraphs]:
        raise RuntimeError("تغير النص بعد حفظ Word وإعادة قراءته.")


def embed_fonts(path, style):
    """Embed licensed fonts using OOXML font obfuscation, preserving original files."""
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(path) as z:
        files = {name: z.read(name) for name in z.namelist()}
    font_table = ET.fromstring(files["word/fontTable.xml"])
    font = next((f for f in font_table if f.attrib.get(qn("w:name")) == style.family), None)
    if font is None:
        font = ET.SubElement(font_table, qn("w:font"), {qn("w:name"): style.family})
    relpath = "word/_rels/fontTable.xml.rels"
    rels = ET.fromstring(files[relpath]) if relpath in files else ET.Element(f"{{{relns}}}Relationships")
    ctypes = ET.fromstring(files["[Content_Types].xml"])
    ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
    if not any(e.attrib.get("Extension") == "odttf" for e in ctypes):
        ET.SubElement(ctypes, f"{{{ctns}}}Default", Extension="odttf",
                      ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
    for kind, filepath in (("Regular", style.regular), ("Bold", style.bold)):
        guid = uuid.uuid5(uuid.NAMESPACE_URL, hashlib.sha256(Path(filepath).read_bytes()).hexdigest())
        key = guid.bytes[::-1]
        data = bytearray(Path(filepath).read_bytes())
        for i in range(32):
            data[i] ^= key[i % 16]
        filename = f"fonts/{guid.hex}.odttf"
        files[f"word/{filename}"] = data
        rid = f"screenplayFont{kind}"
        ET.SubElement(rels, f"{{{relns}}}Relationship", Id=rid,
                      Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                      Target=filename)
        ET.SubElement(font, qn(f"w:embed{kind}"), {
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": rid,
            qn("w:fontKey"): "{" + str(guid).upper() + "}",
        })
    files["word/fontTable.xml"] = ET.tostring(font_table, encoding="utf-8", xml_declaration=True)
    files[relpath] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    files["[Content_Types].xml"] = ET.tostring(ctypes, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for name, content in files.items():
            z.writestr(name, content)


def soffice_binary():
    runtime = os.environ.get("CODEX_PRIMARY_RUNTIME_ROOT")
    candidates = []
    if runtime:
        candidates += [Path(runtime) / "dependencies/bin/override/soffice",
                       Path(runtime) / "bin/override/soffice"]
    candidates += [Path(sys.executable).resolve().parents[2] / "bin/override/soffice"]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if not found:
        raise RuntimeError("LibreOffice غير موجود؛ ثبّته لتصدير PDF.")
    return found


def export_pdf(docx, style):
    import html
    with tempfile.TemporaryDirectory(prefix="screenplay-lo-") as tmp:
        tmp = Path(tmp)
        config = tmp / "fonts.conf"
        config.write_text(
            '<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
            '<include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
            + "".join(f"<dir>{html.escape(str(Path(f).resolve().parent))}</dir>" for f in {style.regular, style.bold})
            + f'<cachedir>{html.escape(str(tmp / "font-cache"))}</cachedir></fontconfig>', encoding="utf-8")
        env = dict(os.environ, FONTCONFIG_FILE=str(config))
        result = subprocess.run([soffice_binary(), f"-env:UserInstallation={(tmp / 'profile').as_uri()}",
                                 "--headless", "--convert-to", "pdf:writer_pdf_Export",
                                 "--outdir", str(docx.parent), str(docx)],
                                env=env, capture_output=True, text=True, timeout=120)
        pdf = docx.with_suffix(".pdf")
        if result.returncode or not pdf.exists() or pdf.stat().st_size < 1000:
            raise RuntimeError(f"فشل تصدير PDF: {(result.stderr or result.stdout)[-1500:]}")
        return pdf


def inspect_pdf(pdf, pages, style):
    import fitz
    errors = []
    fonts = {}
    with fitz.open(pdf) as document:
        if len(document) != len(pages):
            errors.append(f"عدد الصفحات {len(document)} لا يساوي المخطط {len(pages)}.")
        for i, page in enumerate(document):
            if abs(page.rect.width - 210 * MM_PT) > 1 or abs(page.rect.height - 297 * MM_PT) > 1:
                errors.append(f"حجم الورقة غير مطابق في الصفحة {i + 1}.")
            text = page.get_text()
            if "\ufffd" in text or "\u25a1" in text:
                errors.append(f"محرف مفقود محتمل في الصفحة {i + 1}.")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        if span["text"].strip() and (x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1):
                            errors.append(f"نص خارج الورقة في الصفحة {i + 1}.")
            for entry in page.get_fonts(full=True):
                xref, _, _, name, *_ = entry
                blob = document.extract_font(xref)[3]
                fonts[name] = bool(blob)
        if not fonts or not all(fonts.values()):
            errors.append("يوجد خط PDF غير مضمن.")
        family = style.family.replace(" ", "").lower()
        if not any(family in name.replace(" ", "").lower() for name in fonts):
            errors.append("الخط المطلوب غير موجود في PDF.")
        return {"pdf_pages": len(document), "expected_pages": len(pages), "embedded_fonts": fonts,
                "pdf_errors": errors, "pdf_pass": not errors}


def review_source(source):
    from last_line_review import review_dialogue
    reports = []
    for episode in source["episodes"]:
        for scene in episode["scenes"]:
            lines = [f"{b['character']}: {b['dialogue']}" for b in scene["beats"] if b.get("dialogue", "").strip()]
            if lines:
                reports.append({"episode": episode["number"], "scene": scene["number"],
                                "review": review_dialogue("\n".join(lines))})
    return {"engine": "last-line-path", "automated_indicators": 5, "changes_text": False,
            "note": "هذه مؤشرات محدودة وليست اختباراً للقواعد المئة.", "scenes": reports}


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def check_font_coverage(source, style):
    fonts = [TTFont(style.regular), TTFont(style.bold)]
    try:
        available = set.intersection(*(set(f.getBestCmap()) for f in fonts))
        printed = []
        for key in ("title", "part_label", "author", "director"):
            printed.append(source[key])
        printed.extend(source["production"][k] for k in ("organization", "sector", "unit"))
        for episode in source["episodes"]:
            printed += [episode["number"], episode["label"]]
            for scene in episode["scenes"]:
                printed += [scene[k] for k in ("number", "heading", "location", "interior", "time")]
                printed += [b.get(k, "") for b in scene["beats"] for k in FIELDS]
        missing = {ord(c) for c in "".join(printed) if not c.isspace() and unicodedata.category(c) != "Cf" and ord(c) not in available}
        if missing:
            raise ValueError("الخط لا يغطي بعض المحارف المطلوبة: " + ", ".join(f"U+{c:04X}" for c in sorted(missing)[:20]))
    finally:
        for font in fonts:
            font.close()


def format_file(input_path, output_dir, stem="screenplay", style=None, pdf=True, last_line_review=False):
    if not stem or Path(stem).name != stem or stem in (".", ".."):
        raise ValueError("اسم الناتج يجب أن يكون اسماً بسيطاً بلا مسارات.")
    style = (style or Style()).validate()
    source = load_source(input_path)
    check_font_coverage(source, style)
    pages = paginate(source, style)
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    base = output_dir / stem
    write_json(base.with_suffix(".source.json"), source)
    write_json(base.with_suffix(".layout.json"), {"style": asdict(style), "pages": pages})
    docx = base.with_suffix(".docx")
    render_docx(source, pages, style, docx)
    qa = {"source_preserved": True, "docx_text_preserved": True,
          "input_sha256": hashlib.sha256(Path(input_path).read_bytes()).hexdigest(),
          "profile": style.profile, "expected_pages": len(pages), "visual_review": "required",
          "original_font_verified": False, "measurements": "inferred_from_photographs",
          "pdf_text_extraction": "not_guaranteed; use source.json for machine-readable Arabic"}
    if pdf:
        qa.update(inspect_pdf(export_pdf(docx, style), pages, style))
        qa["automated_pass"] = qa["source_preserved"] and qa["pdf_pass"]
    else:
        qa.update(pdf_pass=None, automated_pass=False, note="Word فقط؛ لم يُفحص PDF.")
    write_json(base.with_suffix(".qa.json"), qa)
    if last_line_review:
        write_json(base.with_suffix(".review.json"), review_source(source))
    if pdf and not qa["automated_pass"]:
        raise RuntimeError("فشل فحص الطباعة؛ راجع " + str(base.with_suffix(".qa.json")))
    return {"docx": str(docx), "pdf": str(base.with_suffix(".pdf")) if pdf else None,
            "qa": str(base.with_suffix(".qa.json")), "pages": len(pages), "automated_pass": qa["automated_pass"]}


def main():
    p = argparse.ArgumentParser(description="مسار تنسيق السيناريو العربي")
    p.add_argument("input", type=Path)
    p.add_argument("--out-dir", type=Path, required=True)
    p.add_argument("--stem", default="screenplay")
    p.add_argument("--profile", choices=("reference", "large"), default="reference")
    p.add_argument("--body-size", type=float)
    p.add_argument("--line-height", type=float)
    p.add_argument("--font-file")
    p.add_argument("--bold-font-file")
    p.add_argument("--font-family")
    p.add_argument("--docx-only", action="store_true")
    p.add_argument("--last-line-review", action="store_true")
    args = p.parse_args()
    style = Style(profile=args.profile)
    if args.profile == "large":
        style.body, style.line, style.gap, style.location, style.meta = 18, 29, 10, 22, 16
    if args.body_size is not None:
        style.body = args.body_size
    if args.line_height is not None:
        style.line = args.line_height
    custom = (args.font_file, args.bold_font_file, args.font_family)
    if any(custom):
        if not all(custom):
            p.error("مرر ملفي الخط العادي والغليظ واسم العائلة معاً.")
        style.regular, style.bold, style.family = custom
    try:
        result = format_file(args.input, args.out_dir, args.stem, style,
                             not args.docx_only, args.last_line_review)
    except (ValueError, RuntimeError) as error:
        p.exit(2, str(error) + "\n")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
